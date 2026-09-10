from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Iterable

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import ChannelDut, ChannelRunResult, CurveData, Metrics, SweepSettings


def application_data_dir() -> Path:
    override = os.environ.get("PCE_STUDIO_DATA_DIR")
    if override:
        return Path(override)
    root = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA")
    if root:
        return Path(root) / "PCEStudio"
    return Path.home() / ".pce_studio"


def default_output_dir() -> Path:
    return Path.home() / "Documents" / "PCE_Data"


def _safe_filename_component(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value.strip())
    value = value.rstrip(". ")
    return value or "unnamed"


def unique_raw_path(output_dir: Path, dut: ChannelDut) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    device = _safe_filename_component(dut.device_id)
    point = _safe_filename_component(dut.point)
    base = output_dir / f"{device}_{dut.condition}_{point}.csv"
    if not base.exists():
        return base
    counter = 2
    while True:
        candidate = output_dir / f"{device}_{dut.condition}_{point}_{counter:03d}.csv"
        if not candidate.exists():
            return candidate
        counter += 1


def write_raw_csv(path: Path, result: ChannelRunResult) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        handle.write("# Software,PCE Studio\n")
        handle.write(f"# Device ID,{result.dut.device_id}\n")
        handle.write(f"# Point,{result.dut.point}\n")
        handle.write(f"# Channel,CH{result.dut.channel}\n")
        handle.write(f"# Condition,{result.dut.condition}\n")
        handle.write(f"# Area (cm^2),{result.dut.area_cm2:.12g}\n")
        handle.write(
            f"# Irradiance (mW/cm^2),{result.settings.irradiance_mw_cm2:.12g}\n"
        )
        writer = csv.writer(handle)
        writer.writerow(
            [
                "Sequence",
                "Direction",
                "Source Mode",
                "Source Value",
                "Voltage (V)",
                "Current (A)",
                "Current Density (mA/cm^2)",
                "Elapsed (s)",
            ]
        )
        sequence = 1
        for direction in result.settings.directions:
            curve = result.curves[direction]
            source_values = result.settings.levels(direction)
            density = curve.current_density(result.dut.area_cm2)
            for source, voltage, current, current_density, elapsed in zip(
                source_values,
                curve.voltage_v,
                curve.current_a,
                density,
                curve.elapsed_s,
                strict=True,
            ):
                writer.writerow(
                    [
                        sequence,
                        direction,
                        result.settings.source_mode,
                        f"{source:.12g}",
                        f"{voltage:.12g}",
                        f"{current:.12g}",
                        f"{current_density:.12g}",
                        f"{elapsed:.12g}",
                    ]
                )
                sequence += 1


class Database:
    def __init__(self, path: Path | None = None) -> None:
        data_dir = application_data_dir()
        data_dir.mkdir(parents=True, exist_ok=True)
        self.path = path or data_dir / "pce_studio.sqlite3"
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(self.path)
        self.connection.row_factory = sqlite3.Row
        self._initialize()

    def close(self) -> None:
        self.connection.close()

    def _initialize(self) -> None:
        self.connection.executescript(
            """
            PRAGMA foreign_keys = ON;
            CREATE TABLE IF NOT EXISTS devices (
                device_id TEXT PRIMARY KEY,
                area_cm2 REAL NOT NULL CHECK(area_cm2 > 0),
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                channel INTEGER NOT NULL,
                device_id TEXT NOT NULL,
                point_name TEXT NOT NULL,
                area_cm2 REAL NOT NULL,
                condition TEXT NOT NULL,
                scan_plan TEXT NOT NULL,
                raw_file TEXT NOT NULL,
                settings_json TEXT NOT NULL,
                metrics_json TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS samples (
                run_id INTEGER NOT NULL REFERENCES runs(id) ON DELETE CASCADE,
                sequence INTEGER NOT NULL,
                direction TEXT NOT NULL,
                voltage_v REAL NOT NULL,
                current_a REAL NOT NULL,
                elapsed_s REAL NOT NULL,
                PRIMARY KEY(run_id, sequence)
            );
            """
        )
        self.connection.commit()

    def save_device(self, device_id: str, area_cm2: float) -> None:
        device_id = device_id.strip()
        if not device_id or area_cm2 <= 0:
            raise ValueError("器件编号不能为空，面积必须大于 0。")
        self.connection.execute(
            """
            INSERT INTO devices(device_id, area_cm2, updated_at)
            VALUES (?, ?, ?)
            ON CONFLICT(device_id) DO UPDATE SET
                area_cm2=excluded.area_cm2,
                updated_at=excluded.updated_at
            """,
            (device_id, area_cm2, datetime.now().isoformat(timespec="seconds")),
        )
        self.connection.commit()

    def delete_device(self, device_id: str) -> None:
        self.connection.execute("DELETE FROM devices WHERE device_id=?", (device_id,))
        self.connection.commit()

    def list_devices(self) -> list[tuple[str, float]]:
        rows = self.connection.execute(
            "SELECT device_id, area_cm2 FROM devices ORDER BY device_id COLLATE NOCASE"
        ).fetchall()
        return [(str(row["device_id"]), float(row["area_cm2"])) for row in rows]

    def save_run(self, result: ChannelRunResult) -> int:
        metric_payload = {
            direction: asdict(metrics) for direction, metrics in result.metrics.items()
        }
        metric_payload["average"] = (
            asdict(result.average_metrics) if result.average_metrics is not None else None
        )
        cursor = self.connection.execute(
            """
            INSERT INTO runs(
                created_at, channel, device_id, point_name, area_cm2, condition,
                scan_plan, raw_file, settings_json, metrics_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                datetime.now().isoformat(timespec="seconds"),
                result.dut.channel,
                result.dut.device_id,
                result.dut.point,
                result.dut.area_cm2,
                result.dut.condition,
                result.settings.scan_plan,
                result.raw_file,
                json.dumps(asdict(result.settings), ensure_ascii=False),
                json.dumps(metric_payload, ensure_ascii=False, allow_nan=False),
            ),
        )
        run_id = int(cursor.lastrowid)
        sequence = 1
        sample_rows: list[tuple[int, int, str, float, float, float]] = []
        for direction in result.settings.directions:
            curve = result.curves[direction]
            for voltage, current, elapsed in zip(
                curve.voltage_v, curve.current_a, curve.elapsed_s, strict=True
            ):
                sample_rows.append(
                    (
                        run_id,
                        sequence,
                        direction,
                        float(voltage),
                        float(current),
                        float(elapsed),
                    )
                )
                sequence += 1
        self.connection.executemany(
            """
            INSERT INTO samples(run_id, sequence, direction, voltage_v, current_a, elapsed_s)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            sample_rows,
        )
        self.connection.commit()
        result.run_id = run_id
        return run_id

    def list_runs(self, limit: int = 200) -> list[dict]:
        rows = self.connection.execute(
            "SELECT * FROM runs ORDER BY id DESC LIMIT ?", (limit,)
        ).fetchall()
        return [dict(row) for row in rows]

    def load_run(self, run_id: int) -> ChannelRunResult:
        row = self.connection.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone()
        if row is None:
            raise KeyError(f"找不到测试记录 {run_id}")
        settings = SweepSettings(**json.loads(row["settings_json"]))
        dut = ChannelDut(
            channel=int(row["channel"]),
            device_id=str(row["device_id"]),
            point=str(row["point_name"]),
            area_cm2=float(row["area_cm2"]),
            condition=str(row["condition"]),
        )
        result = ChannelRunResult(dut=dut, settings=settings, raw_file=str(row["raw_file"]))
        samples = self.connection.execute(
            "SELECT * FROM samples WHERE run_id=? ORDER BY sequence", (run_id,)
        ).fetchall()
        for direction in settings.directions:
            selected = [sample for sample in samples if sample["direction"] == direction]
            result.curves[direction] = CurveData(
                direction=direction,
                voltage_v=np.asarray([sample["voltage_v"] for sample in selected], dtype=float),
                current_a=np.asarray([sample["current_a"] for sample in selected], dtype=float),
                elapsed_s=np.asarray([sample["elapsed_s"] for sample in selected], dtype=float),
            )
        metric_payload = json.loads(row["metrics_json"])
        for direction in settings.directions:
            if metric_payload.get(direction):
                result.metrics[direction] = Metrics(**metric_payload[direction])
        if metric_payload.get("average"):
            result.average_metrics = Metrics(**metric_payload["average"])
        result.run_id = run_id
        return result

    def export_summary(self, run_ids: Iterable[int], output_path: Path) -> Path:
        run_ids = list(dict.fromkeys(int(value) for value in run_ids))
        if not run_ids:
            raise ValueError("当前会话没有可导出的测试结果。")
        placeholders = ",".join("?" for _ in run_ids)
        rows = self.connection.execute(
            f"SELECT * FROM runs WHERE id IN ({placeholders}) ORDER BY id", run_ids
        ).fetchall()

        headers = [
            "File Name",
            "Test Time",
            "Channel",
            "Device ID",
            "Point",
            "Light/Dark",
            "Area (cm^2)",
            "Scan Plan",
            "Forward PCE (%)",
            "Forward Voc (V)",
            "Forward Jsc (mA/cm^2)",
            "Forward FF (%)",
            "Forward Pmax (mW/cm^2)",
            "Reverse PCE (%)",
            "Reverse Voc (V)",
            "Reverse Jsc (mA/cm^2)",
            "Reverse FF (%)",
            "Reverse Pmax (mW/cm^2)",
            "Average PCE (%)",
            "Average Voc (V)",
            "Average Jsc (mA/cm^2)",
            "Average FF (%)",
            "Average Pmax (mW/cm^2)",
            "Raw File",
        ]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Results"
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B78")

        def values(metrics: dict | None) -> list[float | None]:
            if not metrics:
                return [None, None, None, None, None]
            return [
                metrics["pce_percent"],
                metrics["voc_v"],
                metrics["jsc_ma_cm2"],
                metrics["ff_percent"],
                metrics["pmax_mw_cm2"],
            ]

        for row in rows:
            metrics = json.loads(row["metrics_json"])
            raw_path = Path(str(row["raw_file"]))
            sheet.append(
                [
                    raw_path.name,
                    row["created_at"],
                    f"CH{row['channel']}",
                    row["device_id"],
                    row["point_name"],
                    row["condition"],
                    row["area_cm2"],
                    row["scan_plan"],
                    *values(metrics.get("forward")),
                    *values(metrics.get("reverse")),
                    *values(metrics.get("average")),
                    str(raw_path),
                ]
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(sheet.columns, start=1):
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[get_column_letter(index)].width = min(max(width, 12), 42)
        for row in sheet.iter_rows(min_row=2, min_col=7, max_col=23):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path


def timestamped_summary_path(output_dir: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_dir / f"output_file_{stamp}.xlsx"
