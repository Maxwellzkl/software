import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from pce_app.models import ChannelDut, ChannelRunResult, CurveData, Metrics, SweepSettings
from pce_app.storage import Database, unique_raw_path, write_raw_csv


def _result() -> ChannelRunResult:
    settings = SweepSettings(start=-0.1, stop=1.0, count=3, scan_plan="both")
    dut = ChannelDut(1, "D01", "3", 0.04, "light")
    result = ChannelRunResult(dut=dut, settings=settings)
    for direction, voltage in (
        ("forward", np.array([-0.1, 0.45, 1.0])),
        ("reverse", np.array([1.0, 0.45, -0.1])),
    ):
        result.curves[direction] = CurveData(
            direction,
            voltage,
            np.array([-0.0008, -0.0004, 0.0001])[:: 1 if direction == "forward" else -1],
            np.array([0.0, 0.1, 0.2]),
        )
        result.metrics[direction] = Metrics(10.0, 0.9, 20.0, 55.0, 10.0)
    result.average_metrics = Metrics(10.0, 0.9, 20.0, 55.0, 10.0)
    return result


def test_device_and_run_roundtrip(tmp_path: Path) -> None:
    database = Database(tmp_path / "test.sqlite3")
    database.save_device("D01", 0.04)
    assert database.list_devices() == [("D01", 0.04)]
    result = _result()
    raw_path = unique_raw_path(tmp_path, result.dut)
    result.raw_file = str(raw_path)
    write_raw_csv(raw_path, result)
    run_id = database.save_run(result)
    loaded = database.load_run(run_id)
    assert loaded.dut.device_id == "D01"
    assert loaded.average_metrics is not None
    assert loaded.average_metrics.pce_percent == 10.0
    assert raw_path.read_text(encoding="utf-8-sig").startswith("# Software,PCE Studio")
    database.close()


def test_duplicate_raw_names_get_counter(tmp_path: Path) -> None:
    dut = ChannelDut(1, "D01", "3", 0.04, "light")
    first = unique_raw_path(tmp_path, dut)
    first.touch()
    second = unique_raw_path(tmp_path, dut)
    assert second.name == "D01_light_3_002.csv"


def test_excel_summary(tmp_path: Path) -> None:
    database = Database(tmp_path / "test.sqlite3")
    result = _result()
    result.raw_file = str(tmp_path / "D01_light_3.csv")
    write_raw_csv(Path(result.raw_file), result)
    run_id = database.save_run(result)
    output = tmp_path / "output_file.xlsx"
    database.export_summary([run_id], output)
    workbook = load_workbook(output, read_only=True)
    sheet = workbook["Results"]
    assert sheet.max_row == 2
    assert sheet["A2"].value == "D01_light_3.csv"
    assert sheet["S2"].value == 10.0
    database.close()
